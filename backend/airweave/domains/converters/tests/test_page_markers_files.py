"""Page markers from real PDF and PowerPoint files (Airweave A4).

The files are generated in the test (PyMuPDF, python-pptx) with invented content.
"""

import pytest

from airweave.domains.converters.page_markers import page_marker, strip_page_markers
from airweave.domains.converters.pdf import PdfConverter
from airweave.domains.converters.text_extractors.pdf import extract_pdf_text, text_to_markdown
from airweave.domains.converters.text_extractors.pptx import _extract_slide, extract_pptx_text

PDF_PAGES = [
    "Acme Corp quarterly report. Revenue rose 12% in the third quarter of the year.",
    "Globex segment results. Operating margin fell to 9% on higher freight costs.",
    "Initech outlook. Management expects backlog to recover in the coming quarter.",
]


@pytest.fixture
def pdf_path(tmp_path):
    import fitz

    doc = fitz.open()
    for text in PDF_PAGES:
        page = doc.new_page()
        page.insert_textbox(fitz.Rect(72, 72, 540, 720), text, fontsize=11)
    path = tmp_path / "acme-q3.pdf"
    doc.save(path)
    doc.close()
    return str(path)


@pytest.fixture
def pptx_path(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for title, body in [
        ("Acme strategy", "Expand into Globex markets during the next fiscal year."),
        ("Initech review", "Integration costs came in below plan for the quarter."),
    ]:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2))
        box.text_frame.text = body
    path = tmp_path / "acme-strategy.pptx"
    prs.save(path)
    return str(path)


class TestPdf:
    async def test_converter_marks_every_page(self, pdf_path):
        results = await PdfConverter(ocr_provider=None).convert_batch([pdf_path])
        clean, starts = strip_page_markers(results[pdf_path])
        assert [p for _, p in starts] == [1, 2, 3]
        for (offset, _), text in zip(starts, PDF_PAGES, strict=True):
            assert clean[offset:].startswith(text.split(".")[0])

    async def test_without_markers_the_text_is_upstreams(self, pdf_path):
        results = await PdfConverter(ocr_provider=None).convert_batch([pdf_path])
        extraction = await extract_pdf_text(pdf_path)
        assert strip_page_markers(results[pdf_path])[0] == text_to_markdown(extraction.full_text)


class TestPptx:
    async def test_slides_marked_before_headings(self, pptx_path):
        text = await extract_pptx_text(pptx_path)
        assert text.startswith(page_marker(1) + "## Slide 1")
        assert page_marker(2) + "## Slide 2" in text
        clean, starts = strip_page_markers(text)
        assert [p for _, p in starts] == [1, 2]
        assert clean[starts[1][0] :].startswith("## Slide 2")

    async def test_without_markers_the_text_is_upstreams(self, pptx_path):
        from pptx import Presentation

        prs = Presentation(pptx_path)
        upstream = "\n\n---\n\n".join(
            _extract_slide(slide, idx) for idx, slide in enumerate(prs.slides, start=1)
        )
        assert strip_page_markers(await extract_pptx_text(pptx_path))[0] == upstream

    async def test_too_little_text_still_returns_none(self, tmp_path):
        """The 50-character check counts the text without markers, as upstream.

        Three blank slides give 44 characters without markers (under 50, so None upstream)
        and 71 with them.
        """
        from pptx import Presentation

        prs = Presentation()
        for _ in range(3):
            prs.slides.add_slide(prs.slide_layouts[6])  # blank: only "## Slide N"
        path = tmp_path / "blank.pptx"
        prs.save(path)
        assert await extract_pptx_text(str(path)) is None
